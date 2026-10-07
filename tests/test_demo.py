"""Angka akhir demo.py diperiksa terhadap data skenarionya, dihitung terpisah dari kode aplikasi."""
import sqlite3

import demo
from openpyxl import load_workbook


def test_demo_numbers(tmp_path, monkeypatch):
    monkeypatch.delenv("FINANCE_NOW", raising=False)
    r = demo.run_demo(tmp_path)

    expenses = sum(a for _, _, a, _ in demo.EXPENSES)
    bills = sum(a for _, a, _ in demo.BILLS) * 2
    fees = demo.WITHDRAW[1] * 2
    expense_total = expenses + bills + fees + demo.PAID_FOR[1]
    income_total = demo.SALARY * 2
    debt_flow = demo.DEBT_BUDI[1] - demo.PAY_BUDI[1] - demo.DEBT_ANDI[1] + demo.PAY_ANDI[1] + demo.PAID_FOR[1]
    deposits = demo.DEPOSIT * 2
    total = sum(demo.OPENING.values()) + income_total - expense_total + debt_flow - deposits
    savings_total = deposits - demo.SAVINGS_SPEND[2]

    bal = r["balance"]["data"]
    assert bal["total_dompet"] == total == 3_076_000
    assert bal["total_budget"] == total and bal["consistent"] is True
    assert bal["total_tabungan"] == savings_total == 350_000
    assert bal["debt_total"] == demo.DEBT_BUDI[1] - demo.PAY_BUDI[1] + demo.PAID_FOR[1]
    assert bal["receivable_total"] == demo.DEBT_ANDI[1] - demo.PAY_ANDI[1]
    assert bal["net_worth"] == total + savings_total + bal["receivable_total"] - bal["debt_total"]

    rep = r["report"]["data"]
    assert rep["income"]["total"] == income_total
    assert rep["expense"]["total"] == expense_total  # belanja dari tabungan tidak ikut
    assert rep["savings_expense"]["total"] == demo.SAVINGS_SPEND[2]

    budgets = {b["name"]: b["balance"] for b in r["budget"]["data"]["budgets"]}
    alloc = dict(demo.ALLOC)
    spent = lambda cat: sum(a for _, _, a, c in demo.EXPENSES if c == cat)  # noqa: E731
    assert budgets["makan"] == alloc["makan"] * 2 - spent("makan") - demo.PAID_FOR[1] + demo.MOVE[3]
    assert budgets["jajan"] == alloc["jajan"] * 2 - spent("jajan") - demo.MOVE[3]
    assert budgets["tempat tinggal"] == 0 and budgets["pulsa & internet"] == 0
    assert budgets["belum teralokasi"] == 692_000  # sama dengan saat tabungan masih budget: tabungan di luar budget
    assert "dana darurat" not in budgets
    sv = r["savings"]["data"]
    assert sv["total"] == savings_total and sv["savings"][0]["percent"] == 11.7

    a = r["analyze"]["data"]
    sep = [x for x in demo.EXPENSES if x[0].startswith("2026-09")]
    sep_total = sum(x[2] for x in sep) + bills // 2 + fees // 2 + demo.PAID_FOR[1]
    assert a["expense_total"] == sep_total and a["income_total"] == demo.SALARY
    assert a["frequent_small"] == [{"note": "kopi", "count": 6, "total": 48_000, "category": "jajan"}]
    assert a["projection"]["projected_expense"] == sep_total  # hari terakhir bulan
    assert len(a["debts"]["due_within_7_days"]) == 2
    assert a["savings"]["expense"]["total"] == demo.SAVINGS_SPEND[2] and a["savings"]["total"] == savings_total

    path = r["export"]["data"]["path"]
    wb = load_workbook(path)
    assert wb.sheetnames == ["Transaksi", "Ringkasan per kategori", "Saldo dompet", "Budget", "Hutang piutang"]
    conn = sqlite3.connect(tmp_path / "finance.db")
    n_tx = conn.execute("SELECT COUNT(*) FROM transactions WHERE deleted_at IS NULL").fetchone()[0]
    conn.close()
    assert wb["Transaksi"].max_row - 1 == n_tx  # semua transaksi demo ada di Agustus-September
