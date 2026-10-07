"""analyze, daily-check, export."""
from datetime import date

import pytest
from openpyxl import load_workbook

from fin.analysis import previous_period
from fin.parse import Period


def P(s, e):
    return Period("x", "x", date.fromisoformat(s), date.fromisoformat(e))


@pytest.mark.parametrize("cur, prev", [
    (("2026-10-01", "2026-10-06"), ("2026-09-01", "2026-09-06")),
    (("2026-09-01", "2026-09-30"), ("2026-08-01", "2026-08-31")),
    (("2026-03-01", "2026-03-31"), ("2026-02-01", "2026-02-28")),
    (("2026-03-01", "2026-03-30"), ("2026-02-01", "2026-02-28")),
    (("2026-01-01", "2026-01-15"), ("2025-12-01", "2025-12-15")),
    (("2026-08-01", "2026-10-06"), ("2026-05-01", "2026-07-06")),
    (("2026-10-05", "2026-10-11"), ("2026-09-28", "2026-10-04")),
])
def test_previous_period(cur, prev):
    p = previous_period(P(*cur))
    assert (p.start.isoformat(), p.end.isoformat()) == prev


def test_analyze_empty_period(fin):
    for args in (["analyze"], ["analyze", "--period", "2026-08"], ["analyze", "--period", "all"],
                 ["analyze", "--period", "today"]):
        r = fin.ok(*args)
        d = r["data"]
        assert d["expense_total"] == 0 and d["savings_rate"] is None
        assert d["daily"]["busiest_day"] is None and d["top_expenses"] == []
        assert d["comparison"]["expense"]["percent"] is None
        assert "Belum ada pemasukan atau pengeluaran" in r["message"]
    fin.err("BAD_PERIOD", "analyze", "--period", "kemarin-lusa")


def test_analyze_facts(wallets):
    fin = wallets
    sep = "2026-09-15 12:00:00"
    fin.ok("add", "--type", "expense", "--item", "makan|100k|makan", now=sep)
    fin.ok("add", "--type", "income", "--item", "gaji|1jt|gaji", now="2026-10-01 08:00:00")
    fin.ok("budget", "alloc", "--item", "makan|50k", now="2026-10-01 09:00:00")
    for day in range(1, 7):
        fin.ok("add", "--type", "expense", "--item", "Kopi |8k|jajan", now=f"2026-10-0{day} 10:00:00")
    fin.ok("add", "--type", "expense", "--item", "kopi|25k|jajan", now="2026-10-02 11:00:00")  # bukan pengeluaran kecil
    fin.ok("add", "--type", "expense", "--item", "nasi padang|60k|makan", now="2026-10-03 12:00:00")
    fin.ok("transfer", "--from", "bri", "--to", "tunai", "--amount", "100k", now="2026-10-04 12:00:00")
    fin.ok("debt", "add", "--direction", "i_owe", "--person", "Budi", "--amount", "70k", "--due", "2026-10-10",
           now="2026-10-04 13:00:00")

    r = fin.ok("analyze", "--period", "this-month")
    a = r["data"]
    assert a["income_total"] == 1_000_000
    assert a["expense_total"] == 48_000 + 25_000 + 60_000 == 133_000
    assert a["net"] == 867_000 and a["savings_rate"] == 86.7
    cats = {c["category"]: c for c in a["expense_by_category"]}
    assert cats["jajan"]["amount"] == 73_000 and cats["jajan"]["count"] == 7
    assert a["comparison"]["previous_period"]["start"] == "2026-09-01"
    assert a["comparison"]["previous_period"]["end"] == "2026-09-06"
    assert a["comparison"]["expense"]["previous"] == 0  # 15 September di luar 1-6 September

    d = a["daily"]
    assert d["days"] == 6 and d["average_per_day"] == 22_167  # 133.000 / 6, dibulatkan
    assert d["busiest_day"] == {"date": "2026-10-03", "amount": 68_000, "count": 2}
    assert d["days_without_spending"] == 0
    assert [t["amount"] for t in a["top_expenses"]] == [60_000, 25_000, 8_000, 8_000, 8_000]
    assert a["frequent_small"] == [{"note": "kopi", "count": 6, "total": 48_000, "category": "jajan"}]

    status = {b["name"]: b for b in a["budget_status"]["budgets"]}
    assert status["makan"]["spent_in_period"] == 60_000 and status["makan"]["balance"] == -10_000
    assert a["budget_status"]["negative"] == ["makan"]
    assert "Budget minus: makan" in r["message"]

    pr = a["projection"]
    assert pr["days_in_month"] == 31 and pr["projected_expense"] == (133_000 * 31 + 3) // 6
    assert a["debts"]["debt_total"] == 70_000 and a["debts"]["due_within_7_days"][0]["days_until_due"] == 4
    assert 5 <= len(r["message"].splitlines()) <= 9

    # bulan lalu: tidak ada proyeksi; dibanding Agustus
    a = fin.ok("analyze", "--period", "2026-09")["data"]
    assert a["projection"] is None and a["expense_total"] == 100_000
    assert a["comparison"]["previous_period"]["text"] == "1–31 Agustus 2026"
    assert a["comparison"]["by_category"][0] == {"category": "makan", "current": 100_000, "previous": 0,
                                                 "difference": 100_000, "percent": None}


def test_analyze_comparison_percent(wallets):
    fin = wallets
    fin.ok("add", "--type", "expense", "--item", "makan|100k|makan", now="2026-08-10 12:00:00")
    fin.ok("add", "--type", "expense", "--item", "makan|150k|makan", now="2026-09-10 12:00:00")
    c = fin.ok("analyze", "--period", "2026-09")["data"]["comparison"]
    assert c["expense"] == {"current": 150_000, "previous": 100_000, "difference": 50_000, "percent": 50.0}
    assert "naik Rp50.000 (50,0%)" in fin.ok("analyze", "--period", "2026-09")["message"]


def test_daily_check_pagi(wallets):
    fin = wallets
    r = fin.ok("daily-check", "--when", "pagi")
    assert r["data"]["send"] is False
    fin.ok("recurring", "add", "kos", "--amount", "500k", "--day", "5")      # lewat 1 hari
    fin.ok("recurring", "add", "wifi", "--amount", "300k", "--day", "9")     # 3 hari lagi
    fin.ok("recurring", "add", "listrik", "--amount", "100k", "--day", "10") # 4 hari lagi: belum
    fin.ok("recurring", "add", "gym", "--amount", "100k", "--day", "1")
    fin.ok("recurring", "pay", "gym")
    fin.ok("debt", "add", "--direction", "i_owe", "--person", "Budi", "--amount", "50k", "--due", "2026-10-08")
    fin.ok("debt", "add", "--direction", "owed_to_me", "--person", "Andi", "--amount", "20k", "--due", "2026-10-20")
    r = fin.ok("daily-check", "--when", "pagi")
    assert r["data"]["send"] is True
    assert [b["name"] for b in r["data"]["bills"]] == ["kos", "wifi"]
    assert [d["person"] for d in r["data"]["debts"]] == ["Budi"]
    assert "kos" in r["message"] and "lewat 1 hari" in r["message"] and "Budi" in r["message"]


def test_daily_check_malam(wallets):
    fin = wallets
    r = fin.ok("daily-check", "--when", "malam", now="2026-10-07 20:00:00")
    assert r["data"]["send"] is True and r["data"]["transactions"] == 0 and "belum ada transaksi" in r["message"]
    fin.ok("add", "--type", "expense", "--item", "makan|20k", "--item", "kopi|8k", now="2026-10-07 12:00:00")
    fin.ok("add", "--type", "income", "--item", "gaji|1jt|gaji", now="2026-10-07 13:00:00")
    r = fin.ok("daily-check", "--when", "malam", now="2026-10-07 20:00:00")
    assert r["data"]["expense_total"] == 28_000 and r["data"]["expense_count"] == 2
    assert "Rp28.000" in r["message"]
    fin.err("BAD_ARGS", "daily-check", "--when", "siang")


def test_export(wallets, tmp_path):
    fin = wallets
    fin.ok("add", "--type", "expense", "--item", "makan|15k|makan", "--item", "kopi|8k")
    fin.ok("debt", "add", "--direction", "i_owe", "--person", "Budi", "--amount", "50k")
    r = fin.ok("export", "--period", "this-month")
    path = r["data"]["path"]
    assert path.endswith(".xlsx") and "exports" in path and r["data"]["period"]["start"] == "2026-10-01"
    wb = load_workbook(path)
    assert wb.sheetnames == ["Transaksi", "Ringkasan per kategori", "Saldo dompet", "Budget", "Hutang piutang"]
    ws = wb["Transaksi"]
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold
    assert ws.max_row - 1 == r["data"]["counts"]["transactions"] == 6  # 3 saldo awal + 2 pengeluaran + hutang
    amounts = [c.value for c in ws["D"][1:]]
    assert all(isinstance(v, int) for v in amounts) and ws["D2"].number_format == "#,##0"
    acc = wb["Saldo dompet"]
    totals = {row[0].value: row[2].value for row in acc.iter_rows(min_row=2) if str(row[0].value).startswith("TOTAL")}
    assert totals == {"TOTAL DOMPET": 700_000 - 23_000 + 50_000, "TOTAL TABUNGAN": 0}
    assert wb["Hutang piutang"]["C2"].value == "Budi" and wb["Hutang piutang"]["F2"].value == 50_000

    out = tmp_path / "sub" / "laporan.xlsx"
    r = fin.ok("export", "--period", "2026-08", "--out", str(out))
    assert r["data"]["path"] == str(out.resolve()) and out.exists()
    assert load_workbook(out)["Transaksi"].max_row == 1  # hanya header
    fin.err("BAD_ARGS", "export", "--out", str(tmp_path / "laporan.csv"))
    fin.err("BAD_PERIOD", "export", "--period", "2027-01")
