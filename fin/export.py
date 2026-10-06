"""Ekspor ke .xlsx. Excel hanya hasil ekspor; sumber kebenaran tetap SQLite."""
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import budgets, debts
from .ledger import ACCOUNT_TYPE_LABEL, TX_SELECT, TYPE_LABEL, balances
from .report import summarize

MONEY = "#,##0"
PCT = "0.0"


def _sheet(wb, title, headers, rows, money_cols=(), pct_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append(list(row))
    for idx, header in enumerate(headers, 1):
        letter = get_column_letter(idx)
        fmt = MONEY if header in money_cols else (PCT if header in pct_cols else None)
        width = len(str(header))
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            if fmt and isinstance(cell.value, int):
                cell.number_format = fmt
            text = f"{cell.value:,}" if isinstance(cell.value, int) and fmt == MONEY else str(cell.value or "")
            width = max(width, len(text))
        ws.column_dimensions[letter].width = min(width + 2, 60)
    return ws


def build(conn, period, path):
    """Tulis workbook ke path. Kembalikan jumlah baris per sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    start, end = period.bounds()

    txs = conn.execute(TX_SELECT + " WHERE t.deleted_at IS NULL AND t.ts BETWEEN ? AND ? ORDER BY t.ts, t.id",
                       (start, end)).fetchall()
    tx_rows = [(r["id"], r["ts"], TYPE_LABEL[r["type"]], r["amount"], r["account"], r["to_account"],
                r["category"], r["budget"], r["note"], r["raw_text"]) for r in txs]
    _sheet(wb, "Transaksi", ["ID", "Waktu", "Jenis", "Jumlah", "Dompet", "Ke dompet", "Kategori", "Budget",
                             "Catatan", "Teks asli"], tx_rows, money_cols={"Jumlah"})

    summary_rows = []
    for kind, label in (("expense", "pengeluaran"), ("income", "pemasukan")):
        sec = summarize(conn, period, kind)
        summary_rows += [(label, c["category"], c["amount"], c["percent"], c["count"]) for c in sec["by_category"]]
        summary_rows.append((label, "TOTAL", sec["total"], 100.0 if sec["total"] else 0.0, sec["count"]))
    _sheet(wb, "Ringkasan per kategori", ["Jenis", "Kategori", "Jumlah", "Persen", "Transaksi"], summary_rows,
           money_cols={"Jumlah"}, pct_cols={"Persen"})

    bals = balances(conn)
    accounts = [r for r in conn.execute("SELECT * FROM accounts ORDER BY id") if not r["archived"] or bals[r["id"]]]
    acc_rows = [(r["name"], ACCOUNT_TYPE_LABEL[r["type"]], bals[r["id"]],
                 "default" if r["is_default"] else ("dihapus" if r["archived"] else "")) for r in accounts]
    acc_rows.append(("TOTAL", "", sum(bals.values()), ""))
    _sheet(wb, "Saldo dompet", ["Dompet", "Tipe", "Saldo", "Keterangan"], acc_rows, money_cols={"Saldo"})

    spent = {r[0]: r[1] for r in conn.execute(
        "SELECT budget_id, SUM(amount) FROM transactions WHERE deleted_at IS NULL AND type = 'expense' "
        "AND ts BETWEEN ? AND ? GROUP BY budget_id", (start, end))}
    ov = budgets.overview(conn)
    bud_rows = [(b["name"], budgets.KIND_LABEL[b["kind"]], b["balance"], spent.get(b["id"], 0),
                 "ditutup" if b["archived"] else "") for b in ov["budgets"]]
    bud_rows.append(("TOTAL", "", ov["total_budget"], sum(spent.values()), ""))
    _sheet(wb, "Budget", ["Budget", "Jenis", "Saldo", "Pengeluaran periode ini", "Keterangan"], bud_rows,
           money_cols={"Saldo", "Pengeluaran periode ini"})

    debt_rows = []
    for d in conn.execute("SELECT * FROM debts WHERE archived = 0 ORDER BY status, id").fetchall():
        i = debts.info(conn, d)
        debt_rows.append((i["id"], debts.DIRECTION_LABEL[i["direction"]], i["person"], i["principal"], i["paid"],
                          i["remaining"], "lunas" if i["status"] == "paid" else "belum lunas", i["due_date"],
                          i["note"]))
    _sheet(wb, "Hutang piutang", ["ID", "Jenis", "Orang", "Pokok", "Dibayar", "Sisa", "Status", "Jatuh tempo",
                                  "Catatan"], debt_rows, money_cols={"Pokok", "Dibayar", "Sisa"})

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"transactions": len(tx_rows), "categories": len(summary_rows), "accounts": len(accounts),
            "budgets": len(ov["budgets"]), "debts": len(debt_rows)}
